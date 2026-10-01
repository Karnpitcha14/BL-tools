"""สร้างใบตรวจนับสต๊อก (P&G Cycle Count) + ชนสต็อก + Report นับซ้ำรอบ 2/3 + สรุปผล

ลำดับงาน (ทำเป็นรายวัน, MC และ FO เดินรอบแยกกันได้)
  1) generate()        : Stock ตั้งต้น + Day  -> ใบนับครั้งที่ 1 (MC, FO)
  2) process_counted() : ใบนับที่กรอก "ยอดตรวจนับ" แล้ว -> เทียบกับยอดตั้งต้นรายบรรทัด
       - บรรทัดไม่ตรง และยังไม่ครบ 3 ครั้ง -> ใบนับครั้งถัดไป (ฟอร์มเดิม เฉพาะบรรทัดที่ไม่ตรง)
       - ชีตใบนับที่ยังไม่กรอกเลย        -> ส่งต่อไปในไฟล์ใหม่ให้กรอกทีหลัง
       - ไม่เหลือใบนับค้าง                 -> สรุปผลการนับประจำวัน (ฉบับสุดท้าย)
ประวัติทุกรอบเก็บในชีตซ่อน _meta / _data ของไฟล์ จึงไม่ต้องมีฐานข้อมูล
"""
import copy
import io
import json
import os
import re
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break, RowBreak

DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
TEMPLATE_PATH = os.path.join(DATA_DIR, "form_template.xlsx")
PLAN_PATH = os.path.join(DATA_DIR, "plan.xlsx")

# ---- ค่าหัวใบนับ (แก้ตรงนี้ได้) ----
SELLER = "22078"
SELLER_NAME = "WLA (P&G)"
ZONE_NAME = "Storage"
STORERKEY = {"MC": "32078", "FO": "42078"}
STATION_DEFAULT = {"MC": "01", "FO": "A0101"}
MAX_ROUNDS = 3

# ---- ผังฟอร์ม (ตามชีต ใบนับสต๊อก เดิม) ----
ROWS_PER_PAGE = 14
FIRST_DATA_ROW = 15
BLOCK = 18
TEMPLATE_LAST_ROW = 32
NCOLS = 14
COUNT_COL = 13  # M = ยอดตรวจนับ

FIELDS = ["id", "zone", "sku", "desc", "loc", "lot", "mfg", "exp", "lpn",
          "locno", "station", "pick", "sys", "c1", "c2", "c3"]


class CountSheetError(Exception):
    """ข้อผิดพลาดที่แสดงให้ผู้ใช้เห็นได้"""


def _sku_text(v):
    """SKU เป็นข้อความเสมอ — กันกรณีพิมพ์ใน Excel แล้วกลายเป็นตัวเลข"""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


# ======================= อ่านข้อมูล =======================
def _read_table(file, header_must_have):
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        header = [str(h).strip() if h is not None else "" for h in header]
        if all(h in header for h in header_must_have):
            out = []
            for r in rows:
                if r is None or all(v is None for v in r):
                    continue
                out.append({header[i]: r[i] for i in range(min(len(header), len(r)))})
            return out
    raise CountSheetError(
        "ไม่พบคอลัมน์ที่ต้องใช้: " + ", ".join(header_must_have) + " — ตรวจว่าอัปโหลดไฟล์ถูกช่อง")


def read_stock_mc(file):
    return [r for r in _read_table(file, ["Sku", "Loc", "Lpn", "Lot", "Qty"]) if r.get("Sku")]


def read_stock_fo(file):
    rows = _read_table(file, ["Seller Item Code", "Sh Qty", "Location No."])
    return [r for r in rows if r.get("Seller Item Code") and r.get("Item Description")]


def load_plan(cycle_file=None):
    plan = {}
    if cycle_file is not None:
        wb = openpyxl.load_workbook(cycle_file, read_only=True, data_only=True)
        if "ชนสต็อก_SKU" not in wb.sheetnames:
            raise CountSheetError("ไฟล์ Cycle ไม่มีชีต ชนสต็อก_SKU")
        for r in wb["ชนสต็อก_SKU"].iter_rows(min_row=4, max_col=4, values_only=True):
            if r[0] is not None and r[3]:
                plan.setdefault(int(r[0]), []).append(_sku_text(r[3]))
    else:
        wb = openpyxl.load_workbook(PLAN_PATH, read_only=True, data_only=True)
        for r in wb.worksheets[0].iter_rows(min_row=2, max_col=3, values_only=True):
            if r[0] is not None and r[2] not in (None, ""):
                plan.setdefault(int(r[0]), []).append(_sku_text(r[2]))
    return plan


# ======================= สร้างรายการนับ =======================
GROUND, UPPER = "ชั้น 1", "ชั้น 2-8"
GROUP_NOTE = {GROUND: "ชั้น 1 (เดินนับ)", UPPER: "ชั้น 2-8 (โฟล์คลิฟท์)"}
_LOC_RE = re.compile(r"^([A-Z]+\d{2})(\d{2})(\d{2})$")   # A453701 = แถว A45 / ล็อค 37 / ชั้น 01


def _mc_loc_parts(loc):
    m = _LOC_RE.match(str(loc or "").strip().upper())
    return (m.group(1), int(m.group(2)), int(m.group(3))) if m else None


def mc_group(loc):
    """ชั้น 1 = เดินนับ, ชั้น 2 ขึ้นไป = ใช้โฟล์คลิฟท์ (Loc รูปแบบอื่น เช่น PICKTO ถือเป็นชั้น 1)"""
    parts = _mc_loc_parts(loc)
    return UPPER if parts and parts[2] >= 2 else GROUND


def _mc_sort_key(r):
    loc, lpn = str(r.get("Loc") or ""), str(r.get("Lpn") or "")
    parts = _mc_loc_parts(loc)
    if not parts:                      # Loc พิเศษ ไว้ท้ายกลุ่มชั้น 1
        return (0, 1, loc, 0, 0, lpn)
    row, bay, level = parts
    return (0 if level == 1 else 1, 0, row, bay, level, lpn)

def build_lines(day, plan, mc_stock, fo_stock):
    skus = plan.get(day)
    if not skus:
        raise CountSheetError(f"แผนนับไม่มีรายการของ Day {day}")
    sku_set = set(skus)
    fo_loc = {_sku_text(r["Seller Item Code"]): r.get("Location No.") or "" for r in fo_stock}

    mc = [r for r in mc_stock if _sku_text(r["Sku"]) in sku_set]
    mc.sort(key=_mc_sort_key)        # ชั้น 1 ก่อน (แถว→ล็อค) แล้วชั้น 2-8 (แถว→ล็อค→ชั้น)
    fo = [r for r in fo_stock if _sku_text(r["Seller Item Code"]) in sku_set]
    fo.sort(key=lambda r: (not r.get("Station"), str(r.get("Station") or ""),
                           str(r.get("Location No.") or "")))   # เรียงตาม Station

    lines = []
    for r in mc:
        sku = _sku_text(r["Sku"])
        lines.append(dict(zone="MC", sku=sku, desc=r.get("Description"), loc=r.get("Loc"),
                          lot=r.get("Lot"), mfg=r.get("Manu Date"), exp=r.get("Exp date"),
                          lpn=r.get("Lpn"), locno=fo_loc.get(sku, ""), station="", pick="",
                          sys=int(r.get("Qty") or 0)))
    for r in fo:
        lines.append(dict(zone="FO", sku=_sku_text(r["Seller Item Code"]), desc=r.get("Item Description"),
                          loc=r.get("Location No."), lot="", mfg=None, exp=None, lpn="",
                          locno=r.get("Location No."), station=r.get("Station"),
                          pick=r.get("PickCode"), sys=int(r.get("Sh Qty") or 0)))
    for i, ln in enumerate(lines):
        ln.update(id=i, c1=None, c2=None, c3=None)
    return lines


# ======================= ฟอร์มใบนับ =======================
def _copy_row(ws, src, dst):
    ws.row_dimensions[dst].height = ws.row_dimensions[src].height
    for c in range(1, NCOLS + 1):
        s, d = ws.cell(src, c), ws.cell(dst, c)
        d._style = copy.copy(s._style)
        d.value = s.value if src > TEMPLATE_LAST_ROW - 4 else None


def _data_row(i):
    p, k = divmod(i, ROWS_PER_PAGE)
    return FIRST_DATA_ROW + BLOCK * p + k


def _add_form(wb, form, zone, lines, meta, round_no, group=None):
    ws = wb.copy_worksheet(form)
    ws.title = f"ใบนับ {zone} {group} ครั้งที่{round_no}" if group else f"ใบนับ {zone} ครั้งที่{round_no}"
    ws.sheet_view.showGridLines = form.sheet_view.showGridLines
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = copy.copy(form.page_margins)
    for r in range(1, 7):
        ws.row_dimensions[r].hidden = True

    pages = max(1, -(-len(lines) // ROWS_PER_PAGE))
    sig_merges = [m for m in ws.merged_cells.ranges if m.min_row > FIRST_DATA_ROW + ROWS_PER_PAGE - 1]
    for p in range(1, pages):
        off = BLOCK * p
        for r in range(FIRST_DATA_ROW, TEMPLATE_LAST_ROW + 1):
            _copy_row(ws, r, r + off)
        for m in sig_merges:
            ws.merge_cells(start_row=m.min_row + off, start_column=m.min_col,
                           end_row=m.max_row + off, end_column=m.max_col)

    ws["C9"] = STORERKEY[zone]
    ws["L9"] = round_no
    ws["C10"] = SELLER
    ws["D10"] = SELLER_NAME
    ws["L10"] = meta["doc_no"]
    ws["C11"] = f"{ZONE_NAME} — {GROUP_NOTE[group]}" if group else ZONE_NAME
    ws["L11"] = STATION_DEFAULT[zone]
    ws["C12"] = f"{meta['date']}  (Day {meta['day']})"
    ws["L12"] = zone

    for i, ln in enumerate(lines):
        r = _data_row(i)
        vals = [i + 1, ln["loc"], ln["sku"], ln["desc"], ln["lot"], ln["mfg"], ln["exp"],
                ln["lpn"], ln["locno"], ln["station"], ln["pick"], ln["sys"]]
        for c, v in enumerate(vals, start=1):
            ws.cell(r, c).value = v if v is not None else ""
        for c in (6, 7):
            ws.cell(r, c).number_format = "dd/mm/yyyy"

    last = TEMPLATE_LAST_ROW + BLOCK * (pages - 1)
    ws.print_area = f"A7:N{last}"
    ws.print_title_rows = "7:14"
    ws.row_breaks = RowBreak()
    for p in range(1, pages):
        ws.row_breaks.append(Break(id=TEMPLATE_LAST_ROW + BLOCK * (p - 1)))
    return ws.title, pages


# ======================= ชีตซ่อน เก็บประวัติ =======================
def _write_state(wb, meta, lines):
    m = wb.create_sheet("_meta")
    m["A1"], m["B1"] = "state", json.dumps(meta, ensure_ascii=False)
    d = wb.create_sheet("_data")
    d.append(FIELDS)
    for ln in lines:
        d.append([ln.get(f) for f in FIELDS])
    m.sheet_state = d.sheet_state = "hidden"


def _read_state(wb):
    if "_meta" not in wb.sheetnames or "_data" not in wb.sheetnames:
        raise CountSheetError("ไฟล์นี้ไม่ใช่ใบนับที่ออกจาก BL Tools (ไม่พบข้อมูลรอบนับ) — "
                              "ใช้ไฟล์ที่ดาวน์โหลดจากหน้านี้แล้วกรอกยอดลงไป")
    meta = json.loads(wb["_meta"]["B1"].value)
    rows = list(wb["_data"].iter_rows(min_row=2, values_only=True))
    lines = [dict(zip(FIELDS, r)) for r in rows]
    return meta, lines


def _norm(v):
    """เทียบค่าในเซลล์แบบไม่สนรูปแบบ (ข้อความ/ตัวเลข/ว่าง)"""
    return "" if v is None else _sku_text(v)


def _parse_qty(v):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    if isinstance(v, str):
        v = v.replace(",", "").strip()
    try:
        f = float(v)
    except (TypeError, ValueError):
        raise ValueError
    if f < 0 or not f.is_integer():
        raise ValueError
    return int(f)


# ======================= สไตล์รายงาน =======================
F = "Tahoma"
THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
OK_FILL = PatternFill("solid", fgColor="E3F1E6")
BAD_FILL = PatternFill("solid", fgColor="FBE3E1")


def _title(ws, text, sub, ncol):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws["A1"] = text
    ws["A1"].font = Font(name=F, size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    ws["A2"] = sub
    ws["A2"].font = Font(name=F, size=10)
    ws["A2"].alignment = Alignment(horizontal="center")


def _header_row(ws, row, headers, widths=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row, c, h)
        cell.font = Font(name=F, size=9, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = 30
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w


def _body_cell(cell, fmt=None, align=None):
    cell.font = Font(name=F, size=9)
    cell.border = BOX
    if fmt:
        cell.number_format = fmt
    cell.alignment = Alignment(horizontal=align, vertical="center",
                               wrap_text=(align == "left"))


def _result_format(ws, rng):
    from openpyxl.formatting.rule import CellIsRule
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ตรง"'], fill=OK_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ไม่ตรง"'], fill=BAD_FILL))


def _print_setup(ws, last_col, last_row, header_rows=None):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f"A1:{last_col}{last_row}"
    if header_rows:
        ws.print_title_rows = header_rows
    ws.page_margins.left = ws.page_margins.right = 0.3


# ======================= คำนวณผล =======================
def _final(ln):
    for k in ("c3", "c2", "c1"):
        if ln.get(k) is not None:
            return ln[k]
    return None


def _times(ln):
    return sum(ln.get(k) is not None for k in ("c1", "c2", "c3"))


def _add_summary_sheet(wb, lines, pending_ids, meta, final):
    ws = wb.create_sheet("สรุปผลการนับ")
    status = "ฉบับสุดท้าย" if final else "ระหว่างนับ — ยังมีใบนับค้างในไฟล์นี้"
    _title(ws, "รายงานสรุปผลการตรวจนับสต๊อก (Cycle Count) — P&G",
           f"SELLER {SELLER} ({SELLER_NAME})  |  ผู้ให้บริการ: BL Fulfillment (Betterland Distribution Center Co., Ltd.)", 14)
    info = [("Day", meta["day"]), ("วันที่นับ", meta["date"]), ("เลขที่ใบตรวจนับ", meta["doc_no"]),
            ("สถานะรายงาน", status)]
    for i, (k, v) in enumerate(info):
        ws.cell(4 + i, 1, k).font = Font(name=F, size=10, bold=True)
        ws.merge_cells(start_row=4 + i, start_column=1, end_row=4 + i, end_column=2)
        ws.cell(4 + i, 3, v).font = Font(name=F, size=10)
        ws.cell(4 + i, 3).alignment = Alignment(horizontal="left")

    skus = {}
    for ln in lines:
        t = skus.setdefault(ln["sku"], dict(desc=ln["desc"], mc_sys=0, fo_sys=0, mc_cnt=0, fo_cnt=0,
                                             mc_open=False, fo_open=False, times=0,
                                             pending=False, uncounted=False, loc_diff=0))
        z = ln["zone"].lower()
        t[f"{z}_sys"] += ln["sys"]
        f = _final(ln)
        if f is None:
            t[f"{z}_open"] = True
            t["uncounted"] = True
        else:
            t[f"{z}_cnt"] += f
            t["loc_diff"] += f != ln["sys"]
        t["times"] = max(t["times"], _times(ln))
        if ln["id"] in pending_ids:
            t["pending"] = True

    T = 23
    hdr = ["ลำดับ", "Sku", "ชื่อสินค้า", "MC ในระบบ", "FO ในระบบ", "รวมในระบบ",
           "MC นับได้", "FO นับได้", "รวมนับได้", "ส่วนต่าง", "ผล", "นับ (ครั้ง)", "สถานะ"]
    _header_row(ws, T, hdr, [6, 17, 44, 11, 11, 11, 11, 11, 11, 10, 9, 8, 24])
    r = T + 1
    for i, (sku, t) in enumerate(skus.items(), 1):
        if t["pending"]:
            st = f"รอนับครั้งที่ {t['times'] + 1}"
        elif t["uncounted"]:
            st = "ยังไม่ได้นับ"
        elif t["mc_cnt"] + t["fo_cnt"] == t["mc_sys"] + t["fo_sys"] and t["loc_diff"]:
            st = f"ยอดรวมตรง แต่ Location ไม่ตรง ({t['loc_diff']} บรรทัด)"
        elif t["mc_cnt"] + t["fo_cnt"] == t["mc_sys"] + t["fo_sys"]:
            st = "ตรงตั้งแต่ครั้งที่ 1" if t["times"] == 1 else f"ตรงหลังนับครั้งที่ {t['times']}"
        else:
            st = f"ไม่ตรง (ยืนยัน {t['times']} ครั้ง)"
        vals = [i, sku, t["desc"], t["mc_sys"], t["fo_sys"], f"=D{r}+E{r}",
                "" if t["mc_open"] else t["mc_cnt"], "" if t["fo_open"] else t["fo_cnt"],
                f'=IF(OR(G{r}="",H{r}=""),"",G{r}+H{r})', f'=IF(I{r}="","",I{r}-F{r})',
                f'=IF(J{r}="","",IF(J{r}=0,"ตรง","ไม่ตรง"))', t["times"] or "", st]
        for c, v in enumerate(vals, 1):
            _body_cell(ws.cell(r, c, v), "#,##0;-#,##0;0" if 4 <= c <= 10 else None,
                       "left" if c in (3, 13) else "center")
        r += 1
    last = r - 1
    rng = lambda col: f"${col}${T + 1}:${col}${last}"
    _result_format(ws, rng("K"))
    sum_row = r
    ws.cell(r, 3, "รวม").font = Font(name=F, size=9, bold=True)
    for c in "DEFGHIJ":
        cell = ws[f"{c}{r}"]
        cell.value = f"=SUM({c}{T + 1}:{c}{last})"
        _body_cell(cell, "#,##0;-#,##0;0", "center")
        cell.font = Font(name=F, size=9, bold=True)

    ws.cell(9, 1, "ภาพรวมผลการนับ (ชนสต็อก MC + FO เทียบยอดตั้งต้น)").font = Font(name=F, size=11, bold=True)
    overview = [
        ("จำนวน SKU ตามแผน", f"=COUNTA({rng('B')})", "0"),
        ("นับครบแล้ว", f'=COUNTIF({rng("K")},"ตรง")+COUNTIF({rng("K")},"ไม่ตรง")', "0"),
        ("ตรงตั้งแต่ครั้งที่ 1", f'=COUNTIF({rng("M")},"ตรงตั้งแต่ครั้งที่ 1")', "0"),
        ("ตรง (ผลล่าสุด)", f'=COUNTIF({rng("K")},"ตรง")', "0"),
        ("ไม่ตรง (ผลล่าสุด)", f'=COUNTIF({rng("K")},"ไม่ตรง")', "0"),
        ("ยอดรวมตรง แต่ Location ไม่ตรง", f'=COUNTIF({rng("M")},"ยอดรวมตรง*")', "0"),
        ("รอนับซ้ำ / ยังไม่ได้นับ", f'=COUNTIF({rng("M")},"รอนับ*")+COUNTIF({rng("M")},"ยังไม่ได้นับ")', "0"),
        ("% ความถูกต้อง ครั้งที่ 1", "=IF(C11=0,0,C12/C11)", "0.0%"),
        ("% ความถูกต้อง ผลล่าสุด", "=IF(C11=0,0,C13/C11)", "0.0%"),
        ("ยอดในระบบ (ชิ้น)", f"=F{sum_row}", "#,##0"),
        ("ยอดนับได้ (ชิ้น)", f"=I{sum_row}", "#,##0"),
        ("ส่วนต่างสุทธิ (ชิ้น)", f"=J{sum_row}", "#,##0;-#,##0;0"),
    ]
    for i, (k, f, fmt) in enumerate(overview):
        row = 10 + i
        ws.cell(row, 1, k).font = Font(name=F, size=9)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell = ws.cell(row, 3, f)
        cell.font = Font(name=F, size=9, bold=True)
        cell.number_format = fmt
        cell.alignment = Alignment(horizontal="left")
    ws.cell(22, 1, "เกิน (+) / ขาด (−)").font = Font(name=F, size=9)
    ws.merge_cells(start_row=22, start_column=1, end_row=22, end_column=2)
    ws.cell(22, 3, f'="เกิน "&TEXT(SUMIF({rng("J")},">0"),"#,##0")&"   ขาด "&TEXT(SUMIF({rng("J")},"<0"),"#,##0")')
    ws.cell(22, 3).font = Font(name=F, size=9, bold=True)

    s = sum_row + 4
    for col, label in ((1, "จัดทำโดย (BL Fulfillment)"), (6, "ตรวจสอบโดย"), (10, "รับทราบ (P&G)")):
        ws.cell(s, col, "( ........................................ )").font = Font(name=F, size=9)
        ws.cell(s + 1, col, label).font = Font(name=F, size=9, bold=True)
        ws.cell(s + 2, col, "วันที่ ........ / ........ / ........").font = Font(name=F, size=9)
    ws.freeze_panes = f"A{T + 1}"
    _print_setup(ws, "M", s + 2, f"{T}:{T}")


def _add_line_sheet(wb, lines, meta):
    ws = wb.create_sheet("บรรทัดที่ Diff")
    _title(ws, "รายละเอียดรายบรรทัดที่ยอดไม่ตรงกับยอดตั้งต้น (ทุกครั้งที่นับ)",
           f"Day {meta['day']}  |  วันที่นับ {meta['date']}  |  เลขที่ {meta['doc_no']}", 13)
    hdr = ["โซน", "Location", "Sku", "ชื่อสินค้า", "LPN", "Lot", "ในระบบ",
           "ครั้งที่ 1", "ครั้งที่ 2", "ครั้งที่ 3", "นับล่าสุด", "ส่วนต่าง", "ผล"]
    _header_row(ws, 4, hdr, [6, 13, 17, 40, 14, 12, 10, 9, 9, 9, 10, 10, 9])
    r = 5
    for ln in lines:
        if ln["c1"] is None or (ln["c1"] == ln["sys"] and ln["c2"] is None):
            continue
        cs = [ln[k] if ln[k] is not None else "" for k in ("c1", "c2", "c3")]
        vals = [ln["zone"], ln["loc"], ln["sku"], ln["desc"], ln["lpn"], ln["lot"], ln["sys"], *cs,
                f'=IF(J{r}<>"",J{r},IF(I{r}<>"",I{r},H{r}))', f"=K{r}-G{r}",
                f'=IF(L{r}=0,"ตรง","ไม่ตรง")']
        for c, v in enumerate(vals, 1):
            _body_cell(ws.cell(r, c, v), "#,##0;-#,##0;0" if 7 <= c <= 12 else None,
                       "left" if c == 4 else "center")
        r += 1
    if r == 5:
        ws.cell(5, 1, "ไม่มีบรรทัดที่ยอดไม่ตรง").font = Font(name=F, size=9)
        r = 6
    else:
        _result_format(ws, f"M5:M{r - 1}")
    ws.freeze_panes = "A5"
    _print_setup(ws, "M", r - 1, "4:4")


# ======================= ประกอบไฟล์ =======================
def _assemble(meta, lines, forms, report):
    """forms = [(zone, round, [ids], group)]  report = ใส่ชีตสรุปหรือไม่"""
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    form = wb["FORM"]
    by_id = {ln["id"]: ln for ln in lines}
    meta = dict(meta, sheets={})
    info = []
    for zone, rnd, ids, group in forms:
        title, pages = _add_form(wb, form, zone, [by_id[i] for i in ids], meta, rnd, group)
        meta["sheets"][title] = {"zone": zone, "round": rnd, "ids": ids, "group": group}
        info.append({"zone": zone, "group": group, "round": rnd, "lines": len(ids), "pages": pages})
    del wb["FORM"]
    if report:
        pending = {i for _, rnd, ids, _g in forms if rnd > 1 for i in ids}
        _add_summary_sheet(wb, lines, pending, meta, final=not forms)
        _add_line_sheet(wb, lines, meta)
    _write_state(wb, meta, lines)
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), info


def generate(day, count_date, mc_file, fo_file, cycle_file=None, round_no=1):
    """ใบนับครั้งที่ 1: คืน (bytes xlsx, info)"""
    if not isinstance(count_date, date):
        raise CountSheetError("วันที่นับไม่ถูกต้อง")
    plan = load_plan(cycle_file)
    lines = build_lines(day, plan, read_stock_mc(mc_file), read_stock_fo(fo_file))
    meta = {"version": 3, "day": day, "date": f"{count_date:%d/%m/%Y}",
            "doc_no": f"PG{count_date:%d%m%Y}"}
    forms = [("MC", 1, [ln["id"] for ln in lines if ln["zone"] == "MC" and mc_group(ln["loc"]) == g], g)
             for g in (GROUND, UPPER)]
    forms.append(("FO", 1, [ln["id"] for ln in lines if ln["zone"] == "FO"], None))
    forms = [f for f in forms if f[2]]
    data, info = _assemble(meta, lines, forms, report=False)
    return data, {"day": day, "date": meta["date"], "forms": info}


def process_counted(file):
    """รับใบนับที่กรอกยอดแล้ว คืน (bytes xlsx, info)"""
    wb = openpyxl.load_workbook(file, data_only=True)
    meta, lines = _read_state(wb)
    if meta.get("version") != 3:
        raise CountSheetError("ไฟล์นี้ออกจากเวอร์ชันเก่า — สร้างใบนับครั้งที่ 1 ใหม่จากหน้านี้")
    if not meta.get("sheets"):
        raise CountSheetError("ไฟล์นี้เป็นสรุปผลฉบับสุดท้ายแล้ว ไม่มีใบนับให้กรอก")

    by_id = {ln["id"]: ln for ln in lines}
    next_forms, counted, missing, bad = [], [], [], []
    for title, s in meta["sheets"].items():
        if title not in wb.sheetnames:
            raise CountSheetError(f"ไม่พบชีต {title} — อย่าเปลี่ยนชื่อหรือลบชีตใบนับ")
        ws, rnd, ids = wb[title], s["round"], s["ids"]
        vals = []
        for i, lid in enumerate(ids):
            row = _data_row(i)
            ln = by_id[lid]
            got = [_norm(ws.cell(row, c).value) for c in (2, 3, 8)]   # LOCATION, รหัสสินค้า, LPN
            want = [_norm(ln["loc"]), _norm(ln["sku"]), _norm(ln["lpn"])]
            if got != want:
                raise CountSheetError(
                    f"{title} ลำดับ {i + 1}: Location/รหัสสินค้า/LPN ไม่ตรงกับที่ออกไป "
                    f"(ควรเป็น {want[0]} / {want[1]}) — อย่าเรียง แทรก ลบ หรือสลับแถว")
            try:
                vals.append(_parse_qty(ws.cell(row, COUNT_COL).value))
            except ValueError:
                vals.append(None)
                bad.append(f"{title} ลำดับ {i + 1}")
        if all(v is None for v in vals) and not any(b.startswith(title) for b in bad):
            next_forms.append((s["zone"], rnd, ids, s.get("group")))   # ยังไม่ได้นับ ส่งต่อ
            continue
        missing += [f"{title} ลำดับ {i + 1} ({by_id[lid]['loc']})"
                    for i, (lid, v) in enumerate(zip(ids, vals)) if v is None]
        for lid, v in zip(ids, vals):
            if v is not None:
                by_id[lid][f"c{rnd}"] = v
        diff_ids = [lid for lid in ids if by_id[lid][f"c{rnd}"] != by_id[lid]["sys"]]
        counted.append({"zone": s["zone"], "group": s.get("group"), "round": rnd,
                        "lines": len(ids), "diff": len(diff_ids)})
        if diff_ids and rnd < MAX_ROUNDS:
            next_forms.append((s["zone"], rnd + 1, diff_ids, s.get("group")))

    if bad:
        raise CountSheetError("ยอดตรวจนับต้องเป็นจำนวนเต็มไม่ติดลบ: " + ", ".join(bad[:8])
                              + (f" และอีก {len(bad) - 8} แถว" if len(bad) > 8 else ""))
    if missing:
        raise CountSheetError(f"กรอกยอดตรวจนับไม่ครบ {len(missing)} แถว (ถ้าไม่มีของให้กรอก 0): "
                              + ", ".join(missing[:8]) + (" …" if len(missing) > 8 else ""))
    if not counted:
        raise CountSheetError("ยังไม่ได้กรอกยอดตรวจนับในใบนับชีตไหนเลย")

    next_forms.sort(key=lambda f: (f[0] != "MC", f[3] != GROUND, f[1]))
    data, info = _assemble(meta, lines, next_forms, report=True)
    return data, {"day": meta["day"], "date": meta["date"], "counted": counted, "forms": info}
